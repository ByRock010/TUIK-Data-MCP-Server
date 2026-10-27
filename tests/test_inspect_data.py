"""inspect_tuik_data: observed column profiles from retained file snapshots (#51)."""

import io
import json
from pathlib import Path

import httpx
import pytest
from openpyxl import Workbook

from tuik_mcp.clients.base import HttpClient
from tuik_mcp.config import PORTAL, Settings
from tuik_mcp.models.search import TuikSearchResult
from tuik_mcp.parsers.models import ParsedTable
from tuik_mcp.parsers.profile import InspectLimits, profile_table, shape, type_hint
from tuik_mcp.server import respond
from tuik_mcp.services import files
from tuik_mcp.services.catalog import CatalogService

FIXTURES = Path(__file__).parent / "fixtures"
FILES = FIXTURES / "files"
FLOW = "flow:DF_ALTINSIZ_DTH+V1.0"
ID = "resource:" + "b" * 64
URL = f"{PORTAL}/api/tr/data/downloads?t=i&p=inspect"
TYPES = {
    "csv": "text/csv",
    "xls": "application/vnd.ms-excel",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def service(body, kind, **settings):
    """A catalog whose flow export and file download both return `body`."""
    calls = []

    def handler(request):
        calls.append(request.url.path)
        if request.url.path == "/api/tr/dataflows":
            return httpx.Response(200, content=(FIXTURES / "flows.json").read_bytes())
        return httpx.Response(200, content=body, headers={"content-type": TYPES[kind]})

    http = HttpClient(
        Settings(min_interval_seconds=0, download_interval_seconds=0, attempts=1, **settings),
        httpx.MockTransport(handler),
    )
    catalog = CatalogService(http)
    catalog.remember(TuikSearchResult(id=ID, title="Tablo", resource_type="table", source_url=URL))
    return http, catalog, calls


def workbook(*sheets):
    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets:
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def column(response, name):
    return next(c for c in response["columns"] if c["name"] == name)


async def test_csv_profile_then_exact_filter_query_from_one_snapshot(monkeypatch):
    parses = []
    original = files.parse_source
    monkeypatch.setattr(
        files, "parse_source", lambda *a, **k: parses.append(a[1]) or original(*a, **k)
    )
    http, catalog, calls = service((FIXTURES / "data.csv").read_bytes(), "csv")
    try:
        first = await respond(catalog.inspect(FLOW, max_values=5))
        assert first["status"] == "ok" and first["format"] == first["extraction"] == "csv"
        assert first["snapshot_id"].startswith("sha256:") and first["source_url"].endswith("/csv")
        assert (first["total_rows"], first["rows_examined"], first["complete"]) == (3, 3, True)
        assert first["column_count"] == 16 and not first["columns_truncated"]
        assert first["sheets"] == [] and first["layout"] is None
        value = column(first, "Gözlem Değeri")
        assert (value["type_hint"], value["decimal_separator"]) == ("number", ",")
        assert value["type_hint_source"] == "inferred"
        # Source strings are preserved exactly; decimal commas are never converted.
        assert value["values"][0] == {"value": "98,4622068329301", "count": 1, "truncated": False}
        period = column(first, "Zaman (2015-01)")
        assert period["type_hint"] == "period" and period["distinct_count"] == 3
        assert period["distinct_exact"] and period["values_complete"]
        status = column(first, "Gözlem Durumu")
        assert (status["type_hint"], status["empty"], status["non_empty"]) == ("unknown", 3, 0)
        assert status["values"] == [] and status["distinct_count"] == 0

        again = await respond(catalog.inspect(FLOW, max_values=5))
        assert again == first | {"retrieved_at": again["retrieved_at"]}
        assert again["retrieved_at"] == first["retrieved_at"]
        key, wanted = "Zaman (2015-01)", period["values"][1]["value"]
        rows = await respond(catalog.data(FLOW, filters={key: wanted}))
        assert rows["total_rows"] == 1 and rows["rows"][0][key] == wanted
        assert rows["snapshot_id"] == first["snapshot_id"]
        assert rows["retrieved_at"] == first["retrieved_at"]
        # One download and one profiling parse serve inspection twice plus the query.
        assert calls.count("/api/tr/dataflows/DF_ALTINSIZ_DTH+V1.0/file/csv") == 1
        assert parses == ["csv"]
    finally:
        await http.close()


async def test_real_xls_profile_preserves_layout_notes_and_duplicate_headers():
    http, catalog, calls = service((FILES / "labour-basic.xls").read_bytes(), "xls")
    try:
        data = await respond(catalog.inspect(ID, max_values=3))
        assert data["status"] == "ok" and data["format"] == "xls" and data["layout"] == "table"
        assert data["sheet"] == "Tablo" and data["sheets"] == ["Tablo"]
        assert data["unit"] and data["table_title"] and data["table_notes"]
        assert data["total_rows"] == 260 and data["complete"]
        names = [c["name"] for c in data["columns"]]
        assert len(names) == len(set(names)) == data["column_count"] == 37
        # Duplicate-looking source headers keep their disambiguated exact keys.
        assert "İşgücü Labour force [column 4]" in names
        years = column(data, "Yıllar - Years [column 1]")
        assert years["type_hint"] == "period" and years["distinct_count"] == 22
        workforce = column(data, "İşgücü Labour force [column 4]")
        assert workforce["type_hint"] == "number" and workforce["decimal_separator"] is None
        month = column(data, "column_2")["values"][0]["value"]
        rows = await respond(catalog.data(ID, filters={"column_2": month}, limit=2))
        assert rows["status"] == "ok" and rows["rows"][0]["column_2"] == month
        assert rows["snapshot_id"] == data["snapshot_id"]
        assert len(calls) == 1
    finally:
        await http.close()


async def test_placeholders_and_zero_padded_codes_are_reported_not_reinterpreted():
    http, catalog, _ = service((FILES / "population.xls").read_bytes(), "xls")
    try:
        data = await respond(catalog.inspect(ID, max_values=200))
        marked = [c for c in data["columns"] if c["placeholders"]]
        assert marked and all(c["placeholders"] == {"-": 30} for c in marked)
        assert all(c["type_hint"] == "number" for c in marked)
        assert all(any(v["value"] == "-" for v in c["values"]) for c in marked)
    finally:
        await http.close()
    http, catalog, _ = service((FILES / "cpi-basket.xlsx").read_bytes(), "xlsx")
    try:
        data = await respond(catalog.inspect(ID, columns=["Kod Code"], max_values=200))
        code = data["columns"][0]
        # "01" and "10" coexist: zero-padded identifiers are not called numbers.
        assert code["type_hint"] == "mixed" and data["format"] == "xlsx"
        assert code["values"][0]["value"] == "01" and len(code["values"]) == 200
        assert code["distinct_count"] == 752 and code["distinct_exact"]
        assert not code["values_complete"]
        assert not data["columns_truncated"] and data["column_count"] == 4
    finally:
        await http.close()


async def test_sheet_selection_raw_layout_and_input_errors():
    body = workbook(
        ("İlk", [["Kod", "Ad"], ["01", "Gıda"]]),
        ("Diğer", [["Kod", "Ad"], ["02", "Konut"], ["03", ""]]),
    )
    http, catalog, calls = service(body, "xlsx")
    try:
        data = await respond(catalog.inspect(ID, sheet="Diğer"))
        assert data["sheets"] == ["İlk", "Diğer"] and data["sheet"] == "Diğer"
        assert data["layout"] == "raw" and data["warnings"]
        assert column(data, "column_2")["empty"] == 1
        for kwargs, code in [
            ({"sheet": "Yok"}, "invalid_input"),
            ({"columns": ["column_9"]}, "invalid_columns"),
            ({"max_values": 0}, "invalid_input"),
            ({"max_values": True}, "invalid_input"),
            ({"columns": []}, "invalid_input"),
            ({"columns": ["x" * 201]}, "invalid_input"),
        ]:
            response = await respond(catalog.inspect(ID, **kwargs))
            assert response["code"] == code, kwargs
        assert len(calls) == 1
    finally:
        await http.close()


async def test_csv_sheet_malformed_files_and_non_tables_keep_their_semantics():
    http, catalog, calls = service(b"a;b\n1;2\n", "csv")
    try:
        response = await respond(catalog.inspect(FLOW, sheet="Sheet1"))
        assert response["code"] == "invalid_input"
        assert not any(path.endswith("/file/csv") for path in calls)
    finally:
        await http.close()
    http, catalog, _ = service(b"a;a\n1;2\n", "csv")
    try:
        response = await respond(catalog.inspect(FLOW))
        assert response["code"] == "malformed_response" and response["source_url"].endswith("csv")
    finally:
        await http.close()
    http, catalog, _ = service((FILES / "ict-publication.pdf").read_bytes(), "pdf")
    try:
        response = await respond(catalog.inspect(ID))
        assert response["status"] == "download_available" and response["format"] == "pdf"
        assert response["columns"] == [] and response["snapshot_id"] is None
    finally:
        await http.close()


async def test_bulletin_with_related_datasets_requires_an_explicit_choice():
    payload = json.loads((FIXTURES / "press_dataflows.json").read_text())
    title = payload["data"]["statisticalTables"][0]["title"]
    flows = [{"id": "DF_SATIS_SEKLI_SATIS_DURUMU_V3+V1.0", "name": title, "downloadable": True}]
    calls = []

    def handler(request):
        calls.append(request.url.path)
        if request.url.path == "/api/tr/press/58347":
            return httpx.Response(200, json=payload)
        return httpx.Response(200, json=flows)

    http = HttpClient(Settings(min_interval_seconds=0), httpx.MockTransport(handler))
    try:
        response = await respond(CatalogService(http).inspect("press:58347"))
        assert response["status"] == "related_datasets" and response["related_datasets"]
        assert not any("/file/" in path or "downloads" in path for path in calls)
    finally:
        await http.close()


async def test_empty_tables_and_partial_scans_are_explicit():
    http, catalog, _ = service(b"a;b\n", "csv")
    try:
        data = await respond(catalog.inspect(FLOW))
        assert (data["total_rows"], data["rows_examined"], data["complete"]) == (0, 0, True)
        assert [c["type_hint"] for c in data["columns"]] == ["unknown", "unknown"]
    finally:
        await http.close()
    body = "kod;değer\n" + "".join(f"{n % 7};{n},5\n" for n in range(50))
    http, catalog, _ = service(body.encode(), "csv", inspect_max_rows=20, inspect_max_distinct=10)
    try:
        data = await respond(catalog.inspect(FLOW, max_values=3))
        assert (data["total_rows"], data["rows_examined"], data["complete"]) == (50, 20, False)
        kod, value = column(data, "kod"), column(data, "değer")
        # Partial scans never claim exact domains.
        assert kod["distinct_count"] == 7 and not kod["distinct_exact"]
        assert not kod["values_complete"] and len(kod["values"]) == 3
        assert value["distinct_count"] is None and not value["distinct_exact"]
    finally:
        await http.close()


def table(columns, rows, **extra):
    return ParsedTable(
        columns=columns,
        rows=[dict(zip(columns, row, strict=True)) for row in rows],
        total_rows=len(rows),
        extraction="csv",
        **extra,
    )


def test_output_and_column_budgets_shrink_values_then_columns():
    names = [f"c{i}" for i in range(12)]
    rows = [[f"{name}-{n:03d}-" + "v" * 40 for name in names] for n in range(30)]
    limits = InspectLimits(max_output_bytes=8 * 1024)
    profile = profile_table(table(names, rows), max_values=30, limits=limits)
    assert profile.columns and all(len(c.values) < 30 for c in profile.columns)
    assert any("limited listed values" in w for w in profile.warnings)
    assert all(not c.values_complete for c in profile.columns)
    few = profile_table(table(names, rows[:1]), max_values=5, limits=InspectLimits(max_columns=3))
    assert [c.name for c in few.columns] == ["c0", "c1", "c2"] and few.columns_truncated
    assert any("first 3 of 12" in w for w in few.warnings)
    chosen = profile_table(
        table(names, rows[:1]), max_values=5, limits=InspectLimits(), columns=["c5", "c5"]
    )
    assert [c.name for c in chosen.columns] == ["c5"] and not chosen.columns_truncated


def test_long_values_are_marked_truncated():
    profile = profile_table(table(["a"], [["x" * 250]]), max_values=5, limits=InspectLimits())
    (value,) = profile.columns[0].values
    assert value.truncated and len(value.value) == 200


@pytest.mark.parametrize(
    "values,header,expected",
    [
        (["1", "2", "-3"], "Sayı", ("number", None)),
        (["1.5", "2.25"], "x", ("number", ".")),
        (["1,5", "2,25"], "x", ("number", ",")),
        (["1.234", "5.678"], "x", ("number", "ambiguous")),
        (["1,234"], "x", ("number", "ambiguous")),
        (["1.234", "5.5"], "x", ("number", ".")),
        (["1.5", "2,5"], "x", ("mixed", None)),
        (["12.5%", "3%"], "Oran", ("number", ".")),
        (["2024-01-15", "2023-12-31"], "x", ("date", None)),
        (["2024-02-30"], "x", ("string", None)),
        (["01.02.2024", "15/03/2024"], "x", ("string", None)),
        (["2024-Q1", "2024M02", "2024-S1", "2024-W05", "2024-01"], "x", ("period", None)),
        (["2024", "2025"], "Yıl", ("period", None)),
        (["2024", "2025"], "Ay / Month", ("period", None)),
        (["2024", "2025"], "Ayrıntı", ("number", None)),
        (["2024", "2024-Q1"], "x", ("period", None)),
        (["01", "02"], "Kod", ("string", None)),
        (["Gıda", "5"], "x", ("mixed", None)),
        (["-", "..", "…"], "x", ("unknown", None)),
        (["-", "4"], "x", ("number", None)),
        ([], "x", ("unknown", None)),
    ],
)
def test_type_hints_never_guess_ambiguous_notation(values, header, expected):
    assert type_hint({shape(v) for v in values}, header) == expected


def test_settings_bound_profiling_budgets(monkeypatch):
    monkeypatch.setenv("TUIK_INSPECT_MAX_ROWS", "123")
    settings = Settings.from_env()
    assert InspectLimits.from_settings(settings).max_rows == 123
    for key in ["inspect_max_rows", "inspect_max_columns", "inspect_max_distinct"]:
        with pytest.raises(ValueError):
            Settings(**{key: 0})
    with pytest.raises(ValueError):
        Settings(inspect_max_output_bytes=1024)

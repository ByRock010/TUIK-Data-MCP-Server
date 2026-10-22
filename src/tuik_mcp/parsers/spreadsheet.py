"""Heuristic reader for TÜİK XLS/XLSX statistical tables.

TÜİK tables share a layout: title rows (Turkish, English, `[unit]`), one or more
header rows, optional section rows ("Endeks - Index"), data rows whose row labels
appear only on the first row of a group, and footer lines (source, footnotes).
Repeated side-by-side blocks are separated by empty spacer columns. Values are
returned as strings rendered with each cell's number format, like CSV values.
"""

import io
import re
import sys
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from xml.etree.ElementTree import iterparse

from tuik_mcp.utils.normalization import turkish_casefold

from .errors import InvalidFiltersError, ParseError, ParseLimitError
from .limits import ParseLimits
from .models import ParsedTable

SECTION_LABEL = "Bölüm / Section"
SOURCE_PREFIXES = ("tüik", "turkstat", "kaynak", "source")
DECIMALS = re.compile(r"\.([0#?]+)")


@dataclass
class Cell:
    value: str | float | bool | datetime | date | None = None
    number_format: str = "General"

    @property
    def empty(self) -> bool:
        return self.value is None or (isinstance(self.value, str) and not self.value.strip())

    @property
    def number(self) -> float | None:
        if isinstance(self.value, bool) or not isinstance(self.value, int | float):
            return None
        return float(self.value)

    @property
    def text(self) -> bool:
        return isinstance(self.value, str) and not self.empty


@dataclass
class Sheet:
    name: str
    rows: list[list[Cell]]

    @property
    def width(self) -> int:
        return max((len(row) for row in self.rows), default=0)


def year_like(value: float) -> bool:
    return value.is_integer() and 1900 <= value <= 2100


def render(cell: Cell) -> str:
    """Render a cell as published: number-format precision, no thousands separators."""
    value = cell.value
    if cell.empty:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        midnight = value.time() == datetime.min.time()
        return value.date().isoformat() if midnight else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    number = float(value)
    section = number_format_section(cell.number_format, number)
    if not section or section.lower() == "general" or not re.search(r"[0#?]", section):
        return render_general(number)
    if "E+" in section.upper():
        return render_general(number)
    # Sections format the magnitude; the sign is always rendered as "-", never "(...)".
    magnitude = abs(number) * (100 if "%" in section else 1)
    match = DECIMALS.search(section)
    decimals = len(match.group(1)) if match else 0
    quantum = Decimal(1).scaleb(-decimals)
    rendered = Decimal(repr(magnitude)).quantize(quantum, rounding=ROUND_HALF_UP)
    sign = "-" if number < 0 and rendered != 0 else ""
    return f"{sign}{rendered:f}" + ("%" if "%" in section else "")


def render_general(number: float) -> str:
    if number.is_integer() and abs(number) < 1e15:
        return str(int(number))
    return format(number, ".10g")


def number_format_section(number_format: str, value: float = 0.0) -> str:
    """Section of an Excel number format that applies to `value`, without literals/colors.

    Excel sections are positive;negative;zero; with two sections zero uses the first.
    """
    cleaned = re.sub(r'"[^"]*"|\\.|\[[^\]]*\]|_.|\*.', "", number_format or "")
    sections = cleaned.split(";")
    if value < 0 and len(sections) > 1:
        index = 1
    elif value == 0 and len(sections) > 2:
        index = 2
    else:
        index = 0
    return sections[index].strip()


def header_text(cell: Cell) -> str:
    if cell.empty:
        return ""
    if isinstance(cell.value, str):
        return " ".join(cell.value.split())
    return str(cell.value)


def read_xls(body: bytes, limits: ParseLimits) -> list[Sheet]:
    import xlrd

    book = xlrd.open_workbook(
        file_contents=body,
        formatting_info=True,
        on_demand=True,
        ragged_rows=True,
        logfile=sys.stderr,
    )
    try:
        limits.check(book.nsheets, limits.max_sheets, "Workbook sheets")
        sheets, cells = [], 0
        for index in range(book.nsheets):
            # xlrd materializes this sheet before its dimensions are available;
            # this check bounds our grid, not the library's initial allocation.
            sheet = book.sheet_by_index(index)
            cells += sheet.nrows * sheet.ncols
            limits.check(cells, limits.max_cells, "Workbook cells")
            rows = []
            for r in range(sheet.nrows):
                row = []
                for c in range(sheet.row_len(r)):
                    cell = sheet.cell(r, c)
                    fmt = book.format_map[book.xf_list[cell.xf_index].format_key].format_str
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(cell.value)
                    elif cell.ctype in (xlrd.XL_CELL_NUMBER, xlrd.XL_CELL_TEXT):
                        value = cell.value
                    else:
                        value = None
                    row.append(Cell(value, fmt))
                rows.append(row)
            sheets.append(Sheet(sheet.name, rows))
            book.unload_sheet(index)
        return sheets
    finally:
        book.release_resources()


def check_xlsx(body: bytes, limits: ParseLimits) -> None:
    # Inspect before openpyxl inflates shared strings or creates workbook objects.
    from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        members = archive.infolist()
        limits.check(len(members), 4096, "XLSX archive members")
        limits.check(
            sum(m.file_size for m in members),
            limits.max_uncompressed_bytes,
            "XLSX uncompressed bytes",
        )
        sheets = [m for m in members if re.fullmatch(r"xl/worksheets/[^/]+\.xml", m.filename)]
        limits.check(len(sheets), limits.max_sheets, "Workbook sheets")
        cells = 0
        for member in sheets:
            max_row = max_col = row_number = col_in_row = 0
            with archive.open(member) as stream:
                for event, element in iterparse(stream, events=("start", "end")):
                    tag = element.tag.rsplit("}", 1)[-1]
                    if event == "start" and tag == "row":
                        row_number = int(element.get("r", row_number + 1))
                        col_in_row = 0
                        max_row = max(max_row, row_number)
                    if event == "start" and tag == "c":
                        coordinate = element.get("r")
                        if coordinate:
                            col, row = coordinate_from_string(coordinate)
                            col_in_row = column_index_from_string(col)
                            max_row = max(max_row, row)
                        else:
                            col_in_row += 1
                        max_col = max(max_col, col_in_row)
                    limits.check(
                        cells + max_row * max(1, max_col), limits.max_cells, "Workbook cells"
                    )
                    if event == "end":
                        element.clear()
            cells += max_row * max_col


def read_xlsx(body: bytes, limits: ParseLimits) -> list[Sheet]:
    from openpyxl import load_workbook

    check_xlsx(body, limits)
    book = load_workbook(io.BytesIO(body), read_only=True, data_only=True, keep_links=False)
    try:
        limits.check(len(book.worksheets), limits.max_sheets, "Workbook sheets")
        sheets, count = [], 0
        for sheet in book.worksheets:
            # Dimensions are publisher-controlled. Ignore them and also count actual rows.
            sheet.reset_dimensions()
            rows, width = [], 0
            for row in sheet.iter_rows():
                width = max(width, len(row))
                limits.check(count + (len(rows) + 1) * width, limits.max_cells, "Workbook cells")
                rows.append([Cell(c.value, getattr(c, "number_format", "General")) for c in row])
            count += len(rows) * width
            sheets.append(Sheet(sheet.title, rows))
        return sheets
    finally:
        book.close()


def load_sheets(body: bytes, fmt: str, limits: ParseLimits | None = None) -> list[Sheet]:
    limits = limits or ParseLimits()
    try:
        return read_xls(body, limits) if fmt == "xls" else read_xlsx(body, limits)
    except ParseLimitError:
        raise
    except Exception as exc:
        raise ParseError(f"TÜİK returned an unreadable {fmt.upper()} workbook") from exc


@dataclass
class Layout:
    """Row roles detected in a sheet; indices are zero-based source rows."""

    blocks: list[tuple[int, int]]
    title_rows: list[int]
    annotation_rows: list[int]
    header_rows: list[int]
    first_data_row: int
    footer_start: int


class TableLayout:
    def __init__(self, sheet: Sheet):
        self.sheet = sheet
        self.width = sheet.width
        self.grid = [row + [Cell()] * (self.width - len(row)) for row in sheet.rows]
        empty_columns = {c for c in range(self.width) if all(row[c].empty for row in self.grid)}
        self.blocks: list[tuple[int, int]] = []
        start = None
        for c in range(self.width + 1):
            filled = c < self.width and c not in empty_columns
            if filled and start is None:
                start = c
            elif not filled and start is not None:
                self.blocks.append((start, c))
                start = None

    def block_of(self, column: int) -> int | None:
        return next((i for i, (a, b) in enumerate(self.blocks) if a <= column < b), None)

    def cells(self, r: int) -> list[int]:
        return [c for c, cell in enumerate(self.grid[r]) if not cell.empty]

    def blank(self, r: int) -> bool:
        return not self.cells(r)

    def single_text_per_block(self, r: int, *, at_block_start: bool) -> bool:
        """One text cell per block at most; optionally only at block start columns."""
        cells = self.cells(r)
        if not cells or any(not self.grid[r][c].text for c in cells):
            return False
        blocks = [self.block_of(c) for c in cells]
        if len(set(blocks)) != len(blocks):
            return False
        if at_block_start:
            return all(c == self.blocks[b][0] for c, b in zip(cells, blocks, strict=True))
        return True

    def detect(self) -> Layout | None:
        rows = range(len(self.grid))
        footer_start = len(self.grid)
        for r in reversed(rows):
            if self.blank(r) or self.single_text_per_block(r, at_block_start=True):
                footer_start = r
            else:
                break
        title_rows, annotation_rows, r = [], [], 0
        while r < footer_start:
            if self.blank(r):
                r += 1
            elif self.single_text_per_block(r, at_block_start=True):
                title_rows.append(r)
                r += 1
            elif self.annotation(r):
                annotation_rows.append(r)
                r += 1
            else:
                break
        first_data = next((row for row in range(r, footer_start) if self.data_row(row)), None)
        if first_data is None:
            return None
        pre_data = []
        for row in range(r, first_data):
            if self.blank(row):
                continue
            if self.annotation(row):
                annotation_rows.append(row)
            else:
                pre_data.append(row)
        if not pre_data:
            return None
        header_rows = pre_data
        if len(pre_data) > 1 and self.single_text_per_block(pre_data[-1], at_block_start=False):
            # A lone label right above the data opens the first section ("Endeks - Index").
            header_rows = pre_data[:-1]
            first_data = pre_data[-1]
        return Layout(
            self.blocks, title_rows, annotation_rows, header_rows, first_data, footer_start
        )

    def data_row(self, r: int) -> bool:
        """A row holding measurements rather than header labels.

        Any number outside 1900–2100 is a measurement. Year-range numbers are header
        labels (the population table's "2024 / 2025" row) only when the table's first
        column, which holds row labels, is empty; a filled label column next to
        numbers marks data whose values happen to fall in that range.
        """
        row = self.grid[r]
        numbers = [c for c, cell in enumerate(row) if cell.number is not None]
        if any(not year_like(row[c].number) for c in numbers):
            return True
        label = self.blocks[0][0] if self.blocks else 0
        return not row[label].empty and any(c > label for c in numbers)

    def annotation(self, r: int) -> bool:
        """Parenthesised notes above the header, e.g. "(Bin kişi - Thousand persons)"."""
        if not self.single_text_per_block(r, at_block_start=False):
            return False
        texts = [self.grid[r][c].value.strip() for c in self.cells(r)]
        return all(t.startswith("(") and t.endswith(")") for t in texts)

    def labels(self, header_rows: list[int]) -> list[str]:
        labels = []
        for c in range(self.width):
            parts = [header_text(self.grid[r][c]) for r in header_rows]
            labels.append(" / ".join(p for p in parts if p))
        counts: dict[str, int] = {}
        for label in labels:
            if label:
                counts[label] = counts.get(label, 0) + 1
        return [
            (f"{label} [column {c + 1}]" if counts[label] > 1 else label)
            if label
            else f"column_{c + 1}"
            for c, label in enumerate(labels)
        ]


def unique_labels(labels: list[str]) -> list[str]:
    """Make every label distinct so no cell is overwritten when rows become dicts."""
    taken = set(labels)
    seen: set[str] = set()
    result = []
    for label in labels:
        candidate, n = label, 2
        # A generated name must not reuse another source label either.
        while candidate in seen or (candidate != label and candidate in taken):
            candidate = f"{label} ({n})"
            n += 1
        seen.add(candidate)
        result.append(candidate)
    return result


def label_columns(layout: TableLayout, data_rows: list[int]) -> set[int]:
    """Leading block columns holding row labels (text or years) that may be forward-filled."""
    columns = set()
    for start, end in layout.blocks:
        for c in range(start, min(start + 2, end)):
            values = [layout.grid[r][c] for r in data_rows if not layout.grid[r][c].empty]
            if values and all(
                cell.text or ((n := cell.number) is not None and year_like(n)) for cell in values
            ):
                columns.add(c)
    return columns


def parse_sheet(
    sheet: Sheet, extraction: str, *, limit: int | None, filters: dict[str, str]
) -> ParsedTable:
    table = TableLayout(sheet)
    layout = table.detect()
    if layout is None:
        return raw_table(table, sheet.name, extraction, limit=limit, filters=filters)

    title_texts = [table.grid[r][table.cells(r)[0]].value.strip() for r in layout.title_rows]
    unit = next(
        (t[1:-1].strip() for t in title_texts if t.startswith("[") and t.endswith("]")), None
    )
    titles = [t for t in title_texts if not (t.startswith("[") and t.endswith("]"))]
    columns = table.labels(layout.header_rows)

    footer = [
        table.grid[r][c].value.strip()
        for r in range(layout.footer_start, len(table.grid))
        for c in table.cells(r)
    ]
    annotations = [
        table.grid[r][table.cells(r)[0]].value.strip() for r in sorted(layout.annotation_rows)
    ]
    notes = list(dict.fromkeys(annotations + titles[2:] + footer))
    source_note = next((n for n in footer if turkish_casefold(n).startswith(SOURCE_PREFIXES)), None)

    candidate_rows = [
        r for r in range(layout.first_data_row, layout.footer_start) if not table.blank(r)
    ]
    section_rows = {
        r
        for r in candidate_rows
        if table.single_text_per_block(r, at_block_start=False)
        and not any(table.grid[r][c].number is not None for c in table.cells(r))
    }
    data_rows = [r for r in candidate_rows if r not in section_rows]
    fill_columns = label_columns(table, data_rows)

    section_blocks = sorted(
        {table.block_of(c) for r in section_rows for c in table.cells(r)} - {None}
    )
    section_columns = {
        b: SECTION_LABEL
        if len(section_blocks) == 1
        else f"{SECTION_LABEL} [column {table.blocks[b][0] + 1}]"
        for b in section_blocks
    }
    all_columns = unique_labels(columns + list(section_columns.values()))
    columns = all_columns[: len(columns)]
    section_columns = dict(zip(section_columns, all_columns[len(columns) :], strict=True))
    unknown = set(filters) - set(all_columns)
    if unknown:
        raise InvalidFiltersError("Use exact column names returned by get_tuik_data")

    rows, total = [], 0
    sections: dict[int, str] = {}
    previous: dict[int, str] = {}
    for r in range(layout.first_data_row, layout.footer_start):
        if table.blank(r):
            previous.clear()
            continue
        if r in section_rows:
            for c in table.cells(r):
                block = table.block_of(c)
                sections[block] = table.grid[r][c].value.strip()
                for column in range(*table.blocks[block]):
                    previous.pop(column, None)
            continue
        values = [render(cell) for cell in table.grid[r]]
        for c in sorted(fill_columns):
            if values[c]:
                previous[c] = values[c]
            elif c in previous and any(
                values[k] for k in range(c + 1, table.blocks[table.block_of(c)][1])
            ):
                values[c] = previous[c]
        row = dict(zip(columns, values, strict=True))
        for block, label in section_columns.items():
            row[label] = sections.get(block, "")
        if all(row[key] == value for key, value in filters.items()):
            total += 1
            if limit is None or len(rows) < limit:
                rows.append(row)

    warnings = []
    if len(table.blocks) > 1 and len(section_blocks) > 1:
        warnings.append(
            "Side-by-side blocks share column names; each block's section is in its "
            f"'{SECTION_LABEL} [column N]' column, N being the block's first source column."
        )
    return ParsedTable(
        sheet=sheet.name,
        title=titles[0] if titles else None,
        title_en=titles[1] if len(titles) > 1 else None,
        unit=unit,
        columns=all_columns,
        rows=rows,
        total_rows=total,
        notes=notes,
        source_note=source_note,
        extraction=extraction,
        warnings=warnings,
    )


def raw_table(
    table: TableLayout,
    name: str,
    extraction: str,
    *,
    limit: int | None,
    filters: dict[str, str],
) -> ParsedTable:
    columns = [f"column_{c + 1}" for c in range(table.width)]
    if set(filters) - set(columns):
        raise InvalidFiltersError("Use exact column names returned by get_tuik_data")
    rows, total = [], 0
    for r in range(len(table.grid)):
        if table.blank(r):
            continue
        row = dict(zip(columns, (render(cell) for cell in table.grid[r]), strict=True))
        if all(row[key] == value for key, value in filters.items()):
            total += 1
            if limit is None or len(rows) < limit:
                rows.append(row)
    return ParsedTable(
        sheet=name,
        layout="raw",
        columns=columns,
        rows=rows,
        total_rows=total,
        extraction=extraction,
        warnings=[
            "No header/data layout was recognized; rows are the raw grid with source "
            "column positions as names."
        ],
    )


def parse_spreadsheet(
    body: bytes,
    fmt: str,
    *,
    limit: int | None = None,
    filters: dict[str, str] | None = None,
    sheet: str | None = None,
    budgets: ParseLimits | None = None,
) -> list[ParsedTable]:
    """Parse every sheet, or only the named one, of an XLS/XLSX workbook."""
    sheets = load_sheets(body, fmt, budgets)
    if sheet is not None:
        names = [s.name for s in sheets]
        sheets = [s for s in sheets if s.name == sheet]
        if not sheets:
            raise ValueError(f"Workbook has no sheet named {sheet!r}; sheets: {names}")
    return [parse_sheet(s, fmt, limit=limit, filters=filters or {}) for s in sheets]
